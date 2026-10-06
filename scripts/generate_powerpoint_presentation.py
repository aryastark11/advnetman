#!/usr/bin/env python3
"""
Generates a polished 6-Slide PowerPoint Presentation (.pptx)
Covering Lab 1 to Lab 4/5 Network Automation, Observability & NSoT Framework.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_PPTX = "/home/student/Desktop/lab1/reports/Lab1_to_Lab4_Network_Automation_Presentation.pptx"
os.makedirs(os.path.dirname(OUTPUT_PPTX), exist_ok=True)

# Color Palette (Soothing Modern Dark Theme)
C_BG = RGBColor(15, 23, 42)         # #0F172A Dark Slate
C_CARD = RGBColor(30, 41, 59)       # #1E293B Card Background
C_CARD_BORDER = RGBColor(51, 65, 85) # #334155
C_CYAN = RGBColor(56, 189, 248)     # #38BDF8 Electric Sky Blue (Headings)
C_TEAL = RGBColor(13, 148, 136)     # #0D9488 Teal
C_EMERALD = RGBColor(52, 211, 153)  # #34D399 Emerald (Accents/Pass)
C_AMBER = RGBColor(251, 191, 36)    # #FBBF24 Amber/Gold
C_PURPLE = RGBColor(192, 132, 252)  # #C084FC Purple Accent
C_TEXT_MAIN = RGBColor(241, 245, 249) # #F1F5F9 Slate White
C_TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8 Soft Slate
C_WHITE = RGBColor(255, 255, 255)

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions (13.333" x 7.5")
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide layout

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = C_BG
        bg.line.fill.background()
        return bg

    def add_slide_header(slide, title_text, category_badge="CSCI-5840 ADVANCED NETWORK MANAGEMENT & AUTOMATION", slide_num=1):
        # Header Container
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category Badge
        p_badge = tf.paragraphs[0]
        p_badge.text = category_badge.upper()
        p_badge.font.size = Pt(9.5)
        p_badge.font.bold = True
        p_badge.font.color.rgb = C_EMERALD
        p_badge.space_after = Pt(2)

        # Title
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(20)
        p_title.font.bold = True
        p_title.font.color.rgb = C_CYAN

        # Divider line
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(0.02))
        line.fill.solid()
        line.fill.fore_color.rgb = C_CARD_BORDER
        line.line.fill.background()

        # Footer
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
        ftf = footer_box.text_frame
        ftf.margin_left = ftf.margin_top = ftf.margin_right = ftf.margin_bottom = 0
        p_foot = ftf.paragraphs[0]
        p_foot.text = f"Automated Network Management & Source of Truth Framework  |  Slide {slide_num} of 6"
        p_foot.font.size = Pt(8.5)
        p_foot.font.color.rgb = C_TEXT_MUTED

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Main Title Card Container
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.8), Inches(11.533), Inches(5.9))
    card1.fill.solid()
    card1.fill.fore_color.rgb = C_CARD
    card1.line.color.rgb = C_CARD_BORDER
    card1.line.width = Pt(1.5)

    tb1 = s1.shapes.add_textbox(Inches(1.4), Inches(1.2), Inches(10.5), Inches(5.1))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "CSCI-5840 ADVANCED NETWORK MANAGEMENT & AUTOMATION"
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = C_EMERALD
    p0.space_after = Pt(8)

    p1 = tf1.add_paragraph()
    p1.text = "Enterprise Network Automation & Observability Platform"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = C_CYAN
    p1.space_after = Pt(6)

    p2 = tf1.add_paragraph()
    p2.text = "End-to-End Infrastructure-as-Code, Real-Time Telemetry Data Lake & Network Source of Truth (Labs 1–4)"
    p2.font.size = Pt(13)
    p2.font.color.rgb = C_TEXT_MUTED
    p2.space_after = Pt(22)

    # 3 Pillar Summary Cards inside Slide 1
    pillars = [
        ("Lab 1: Fabric & IPAM", "• Arista cEOS-lab Migration\n• OSPFv2/v3, BGP & RIPv2\n• Dual-Stack DHCPv4/v6 IPAM", C_CYAN),
        ("Lab 2 & 3: Telemetry NOC", "• InfluxDB 7-Day Data Lake\n• Telegraf SNMP & gRPC Ingest\n• Grafana NOC & NetworkX Flow", C_EMERALD),
        ("Lab 4 & 5: NSoT & IaC", "• Django NSoT Web Portal\n• Live eAPI Config Pull & Diff\n• Multi-Vendor Jinja2 (4 OSs)", C_AMBER)
    ]

    for i, (p_title, p_desc, p_color) in enumerate(pillars):
        px = Inches(1.4 + i * 3.55)
        py = Inches(3.2)
        p_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, Inches(3.4), Inches(2.2))
        p_card.fill.solid()
        p_card.fill.fore_color.rgb = C_BG
        p_card.line.color.rgb = p_color
        p_card.line.width = Pt(1.5)

        ptb = s1.shapes.add_textbox(px + Inches(0.15), py + Inches(0.15), Inches(3.1), Inches(1.9))
        ptf = ptb.text_frame
        ptf.word_wrap = True
        ptf.margin_left = ptf.margin_top = ptf.margin_right = ptf.margin_bottom = 0

        p_hdr = ptf.paragraphs[0]
        p_hdr.text = p_title
        p_hdr.font.size = Pt(11)
        p_hdr.font.bold = True
        p_hdr.font.color.rgb = p_color
        p_hdr.space_after = Pt(6)

        p_bod = ptf.add_paragraph()
        p_bod.text = p_desc
        p_bod.font.size = Pt(9.5)
        p_bod.font.color.rgb = C_TEXT_MAIN

    # Presenter metadata footer
    pmeta = s1.shapes.add_textbox(Inches(1.4), Inches(5.8), Inches(10.5), Inches(0.7))
    pmtf = pmeta.text_frame
    pmtf.margin_left = pmtf.margin_top = pmtf.margin_right = pmtf.margin_bottom = 0
    p_meta_text = pmtf.paragraphs[0]
    p_meta_text.text = "Author: Kavyashree Mahadevaiah (aryastark11)   |   Repository: github.com/aryastark11/advnetman"
    p_meta_text.font.size = Pt(10)
    p_meta_text.font.bold = True
    p_meta_text.font.color.rgb = C_EMERALD

    # =========================================================================
    # SLIDE 2: Lab 1 — Multi-Vendor Network Virtualization & IPAM Architecture
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_slide_header(s2, "Lab 1: Multi-Vendor Network Virtualization & IPAM Architecture", 
                     "LAB 1 • TOPOLOGY, PROTOCOLS & IP ADDRESS MANAGEMENT", 2)

    # Left Column: Card Container with bullet points
    c2_left = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    c2_left.fill.solid()
    c2_left.fill.fore_color.rgb = C_CARD
    c2_left.line.color.rgb = C_CARD_BORDER
    c2_left.line.width = Pt(1)

    tb2 = s2.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0

    points_s2 = [
        ("Arista cEOS Fabric Migration:", " Replaced baseline Linux FRR containers with production-grade Arista cEOS (v4.33.10M) across all 5 routers (R1–R5) and 4 switches (S1–S4).", C_CYAN),
        ("Hierarchical Core/Edge Design:", " Dual-chassis Distribution layer (R1-R2), ASBR Core (R3-R4), PE WAN Gateway (R5), Access switches (S1-S2), and Core switches (S3-S4).", C_EMERALD),
        ("Dynamic Routing Protocols:", " OSPFv2/OSPFv3 Area 0 core backbone, iBGP AS 65001 mesh, eBGP AS 65001 -> AS 65002 / AS 65100 peering, and RIPv2 on distribution subinterfaces.", C_AMBER),
        ("Comprehensive Dual-Stack IPAM:", " Structured IPv4/IPv6 IPAM scheme (ipam.csv), dual-stack DHCPv4/DHCPv6 server pools on R2, and SLAAC autoconfiguration for hosts H1–H4.", C_PURPLE),
    ]

    for idx, (title_b, text_b, color_b) in enumerate(points_s2):
        p = tf2.paragraphs[0] if idx == 0 else tf2.add_paragraph()
        p.space_after = Pt(10)
        
        run_t = p.add_run()
        run_t.text = title_b
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = color_b
        
        run_b = p.add_run()
        run_b.text = text_b
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = C_TEXT_MAIN

    # Right Column: Network Topology Diagram Image
    topo_img_path = "/home/student/Desktop/lab1/network_topology.png"
    if os.path.exists(topo_img_path):
        c2_right = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
        c2_right.fill.solid()
        c2_right.fill.fore_color.rgb = C_CARD
        c2_right.line.color.rgb = C_CARD_BORDER
        c2_right.line.width = Pt(1)

        s2.shapes.add_picture(topo_img_path, Inches(6.95), Inches(1.75), Inches(5.433), Inches(4.9))

    # =========================================================================
    # SLIDE 3: Lab 2 — NMAS Telemetry Engine & InfluxDB Data Lake
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_slide_header(s3, "Lab 2: NMAS Telemetry Engine & InfluxDB Data Lake", 
                     "LAB 2 • CENTRALIZED OBSERVABILITY, DATA LAKE & BACKUP RETENTION", 3)

    # Left Column: Points
    c3_left = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    c3_left.fill.solid()
    c3_left.fill.fore_color.rgb = C_CARD
    c3_left.line.color.rgb = C_CARD_BORDER
    c3_left.line.width = Pt(1)

    tb3 = s3.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    tf3 = tb3.text_frame
    tf3.word_wrap = True
    tf3.margin_left = tf3.margin_top = tf3.margin_right = tf3.margin_bottom = 0

    points_s3 = [
        ("Centralized NMAS Station:", " Dedicated Network Management & Automation Station node (172.20.20.100) orchestrating data collection, health checks, and backup pipelines.", C_CYAN),
        ("InfluxDB Time-Series Data Lake:", " High-throughput InfluxDB v1.8 data lake on port 8086 with automated 7-day retention policy (autogen) and automated cleanup of records older than 7 days.", C_EMERALD),
        ("Telegraf Collector Pipeline:", " Automated SNMP polling across 9 cEOS switches/routers (CPU load, memory, interface in/out octets, drop counters) and gRPC streaming telemetry ingest.", C_AMBER),
        ("Isolated Secondary Backup Node:", " Standalone backup server (172.20.20.150) receiving automated data lake snapshots, strictly isolated and accessible exclusively by NMAS.", C_PURPLE),
    ]

    for idx, (title_b, text_b, color_b) in enumerate(points_s3):
        p = tf3.paragraphs[0] if idx == 0 else tf3.add_paragraph()
        p.space_after = Pt(10)
        
        run_t = p.add_run()
        run_t.text = title_b
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = color_b
        
        run_b = p.add_run()
        run_b.text = text_b
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = C_TEXT_MAIN

    # Right Column: Data Lake Diagram Image
    dl_img_path = "/home/student/Desktop/lab1/reports/slide3_datalake_diagram.png"
    if os.path.exists(dl_img_path):
        c3_right = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
        c3_right.fill.solid()
        c3_right.fill.fore_color.rgb = C_CARD
        c3_right.line.color.rgb = C_CARD_BORDER
        c3_right.line.width = Pt(1)

        s3.shapes.add_picture(dl_img_path, Inches(6.95), Inches(1.75), Inches(5.433), Inches(4.9))

    # =========================================================================
    # SLIDE 4: Lab 3 — Real-Time Grafana NOC & Dynamic Traffic Visualization
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_slide_header(s4, "Lab 3: Real-Time Grafana NOC & Dynamic Traffic Visualization", 
                     "LAB 3 • NETWORK OPERATIONS CENTER (NOC) & NETWORKX TRAFFIC SIMULATION", 4)

    # Left Column: Points
    c4_left = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    c4_left.fill.solid()
    c4_left.fill.fore_color.rgb = C_CARD
    c4_left.line.color.rgb = C_CARD_BORDER
    c4_left.line.width = Pt(1)

    tb4 = s4.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    tf4 = tb4.text_frame
    tf4.word_wrap = True
    tf4.margin_left = tf4.margin_top = tf4.margin_right = tf4.margin_bottom = 0

    points_s4 = [
        ("Grafana NOC Master Dashboard:", " Operational on http://localhost:3000/d/nmas-noc-master. Visualizes live SNMP CPU gauge dials, interface throughput (bps), packet error rates, and ICMP ping latency.", C_CYAN),
        ("Dynamic Traffic Flow Simulation:", " Python NetworkX + Matplotlib dynamic simulator rendering animated multi-tier traffic flows (GIF / HTML) with color-coded link utilization thresholds (Green <50%, Yellow 50-80%, Red >80%).", C_EMERALD),
        ("gRPC Streaming Telemetry:", " Continuous sub-second metric streaming directly into the data lake, providing instantaneous fault detection over traditional polling.", C_AMBER),
        ("Bidirectional Navigation:", " Embedded top navigation banner and dashboard link returning directly to the Django NSoT portal on localhost:8000.", C_PURPLE),
    ]

    for idx, (title_b, text_b, color_b) in enumerate(points_s4):
        p = tf4.paragraphs[0] if idx == 0 else tf4.add_paragraph()
        p.space_after = Pt(10)
        
        run_t = p.add_run()
        run_t.text = title_b
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = color_b
        
        run_b = p.add_run()
        run_b.text = text_b
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = C_TEXT_MAIN

    # Right Column: Traffic Flow Snapshot Image
    tf_img_path = "/home/student/Desktop/lab1/reports/traffic_flow_snapshot.png"
    if os.path.exists(tf_img_path):
        c4_right = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
        c4_right.fill.solid()
        c4_right.fill.fore_color.rgb = C_CARD
        c4_right.line.color.rgb = C_CARD_BORDER
        c4_right.line.width = Pt(1)

        s4.shapes.add_picture(tf_img_path, Inches(6.95), Inches(1.75), Inches(5.433), Inches(4.9))

    # =========================================================================
    # SLIDE 5: Lab 4 & 5 — Django Network Source of Truth (NSoT) & Automation Web GUI
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_slide_header(s5, "Lab 4 & 5: Django Network Source of Truth (NSoT) Web Platform", 
                     "LAB 4 & 5 (PART 1) • DJANGO NSOT PORTAL, LIVE eAPI PULL & GOLDEN CONFIGS", 5)

    # Left Column: Points
    c5_left = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    c5_left.fill.solid()
    c5_left.fill.fore_color.rgb = C_CARD
    c5_left.line.color.rgb = C_CARD_BORDER
    c5_left.line.width = Pt(1)

    tb5 = s5.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    tf5 = tb5.text_frame
    tf5.word_wrap = True
    tf5.margin_left = tf5.margin_top = tf5.margin_right = tf5.margin_bottom = 0

    points_s5 = [
        ("Full-Stack Django Portal (Port 8000):", " Single Source of Truth managing 16 network nodes with KPI cards, role filters (Router, Switch, Host, NMAS, Server), and search.", C_CYAN),
        ("Live eAPI Configuration Pull:", " Real-time Arista JSON-RPC eAPI client pulling live running configurations directly from physical/container nodes with batch 'Pull All' capability.", C_EMERALD),
        ("Dynamic Device Provisioning Form:", " Interactive onboarding with vendor profiles, IPAM fields, router routing protocols (OSPF, BGP ASN, Router-ID), switch VLANs, and dual file uploads (.cfg & .j2).", C_AMBER),
        ("Golden Configuration Archive & Diff:", " Timestamped version snapshots (vYYYYMMDD_HHMMSS) stored in SQLite DB and filesystem (golden_configs/) with in-browser Unified Diff comparator.", C_PURPLE),
    ]

    for idx, (title_b, text_b, color_b) in enumerate(points_s5):
        p = tf5.paragraphs[0] if idx == 0 else tf5.add_paragraph()
        p.space_after = Pt(10)
        
        run_t = p.add_run()
        run_t.text = title_b
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = color_b
        
        run_b = p.add_run()
        run_b.text = text_b
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = C_TEXT_MAIN

    # Right Column: NSoT Diagram Image
    nsot_img_path = "/home/student/Desktop/lab1/reports/slide5_nsot_diagram.png"
    if os.path.exists(nsot_img_path):
        c5_right = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
        c5_right.fill.solid()
        c5_right.fill.fore_color.rgb = C_CARD
        c5_right.line.color.rgb = C_CARD_BORDER
        c5_right.line.width = Pt(1)

        s5.shapes.add_picture(nsot_img_path, Inches(6.95), Inches(1.75), Inches(5.433), Inches(4.9))

    # =========================================================================
    # SLIDE 6: Multi-Vendor IaC Templates, Git CI/CD & Future Roadmap
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_slide_header(s6, "Multi-Vendor IaC Templates, Git CI/CD & Future Roadmap", 
                     "MULTI-VENDOR EXTRA CREDIT • ARISTA, CISCO NX-OS/XRv, SONiC & CI/CD", 6)

    # Left Column: Points
    c6_left = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.8), Inches(5.2))
    c6_left.fill.solid()
    c6_left.fill.fore_color.rgb = C_CARD
    c6_left.line.color.rgb = C_CARD_BORDER
    c6_left.line.width = Pt(1)

    tb6 = s6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.4), Inches(4.8))
    tf6 = tb6.text_frame
    tf6.word_wrap = True
    tf6.margin_left = tf6.margin_top = tf6.margin_right = tf6.margin_bottom = 0

    points_s6 = [
        ("Hierarchical Jinja2 Templates (Extra Credit):", " 13 multi-vendor templates across 4 platforms (Arista cEOS, Cisco NX-OS 9000v, Cisco IOS-XRv 9000, SONiC VS) for all 5 network tiers.", C_CYAN),
        ("Decoupled YAML Data Modeling:", " 16 standalone data models in data_models/ separating network business logic from syntax, enabling cross-vendor portability.", C_EMERALD),
        ("Git CI/CD Version Control:", " Production-ready repository live on GitHub (github.com/aryastark11/advnetman) on branch main with automated 6/6 test verification suite.", C_AMBER),
        ("Future Roadmap:", " Next phases: Closed-loop self-healing automation, automated pre-change validation with Batfish, and gNMI telemetry streaming.", C_PURPLE),
    ]

    for idx, (title_b, text_b, color_b) in enumerate(points_s6):
        p = tf6.paragraphs[0] if idx == 0 else tf6.add_paragraph()
        p.space_after = Pt(10)
        
        run_t = p.add_run()
        run_t.text = title_b
        run_t.font.size = Pt(11)
        run_t.font.bold = True
        run_t.font.color.rgb = color_b
        
        run_b = p.add_run()
        run_b.text = text_b
        run_b.font.size = Pt(10)
        run_b.font.color.rgb = C_TEXT_MAIN

    # Right Column: Multi-Vendor Matrix Image
    mv_img_path = "/home/student/Desktop/lab1/reports/slide6_multivendor_matrix.png"
    if os.path.exists(mv_img_path):
        c6_right = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.2))
        c6_right.fill.solid()
        c6_right.fill.fore_color.rgb = C_CARD
        c6_right.line.color.rgb = C_CARD_BORDER
        c6_right.line.width = Pt(1)

        s6.shapes.add_picture(mv_img_path, Inches(6.95), Inches(1.75), Inches(5.433), Inches(4.9))

    # Save presentation
    prs.save(OUTPUT_PPTX)
    print("Successfully generated PowerPoint presentation:", OUTPUT_PPTX)

if __name__ == '__main__':
    create_presentation()
